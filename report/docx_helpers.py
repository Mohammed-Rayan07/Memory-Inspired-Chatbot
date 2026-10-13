"""Small python-docx helpers: headings, paragraphs, bullet lists, auto-numbered tables/figures, page numbers."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(1.0)
        st = self.d.styles["Normal"]
        st.font.name = "Times New Roman"; st.font.size = Pt(11.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
        for name, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 12)):
            h = self.d.styles[name]; h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x2D, 0x4D)
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        self.nt = 0; self.nf = 0
        self._footer()

    def _footer(self):
        p = self.d.sections[0].footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            r._r.append(e)

    def h(self, text, lvl=1):
        return self.d.add_heading(text, lvl)

    def p(self, text="", bold=False, italic=False, align=None, size=None, after=None):
        para = self.d.add_paragraph()
        self._runs(para, text, bold, italic, size)
        if align == "center": para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if align == "justify": para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if after is not None: para.paragraph_format.space_after = Pt(after)
        return para

    def _runs(self, para, text, bold=False, italic=False, size=None):
        """Supports **bold** and *italic* inline markup."""
        import re
        text = text.replace("`", "")
        for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not tok: continue
            if tok.startswith("**"): r = para.add_run(tok[2:-2]); r.bold = True
            elif tok.startswith("*"): r = para.add_run(tok[1:-1]); r.italic = True
            else: r = para.add_run(tok); r.bold = bold
            if italic: r.italic = True
            if size: r.font.size = Pt(size)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style); self._runs(para, it); para.paragraph_format.space_after = Pt(2)

    def numbered(self, items):
        self.bullets(items, "List Number")

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            para = self.d.add_paragraph(); r = para.add_run(line); r.font.name = "Courier New"; r.font.size = Pt(9)
            para.paragraph_format.space_after = Pt(0); para.paragraph_format.left_indent = Inches(0.25)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, header, rows, caption, widths=None, size=9.5, bold_last_col=False):
        self.nt += 1
        cap = self.p(f"**Table {self.nt}.** {caption}", size=10.5, after=3); cap.paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(str(htxt)); r.bold = True; r.font.size = Pt(size)
            tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); tcPr.append(sh)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""; para = cells[i].paragraphs[0]; self._runs(para, str(v), size=size)
                para.paragraph_format.space_after = Pt(0)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths): row.cells[i].width = Inches(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.nt

    def figure(self, path, caption, width=6.2):
        self.nf += 1
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Inches(width)); para.paragraph_format.keep_with_next = True
        self.p(f"**Figure {self.nf}.** {caption}", size=10.5, align="center")
        return self.nf

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.d.save(path)
