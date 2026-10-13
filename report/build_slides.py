"""Build the presentation (pptx) from results.json.  Run: python report/build_slides.py"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
R = json.load(open(ROOT / "results" / "results.json")); FIG = ROOT / "report" / "figures"
M = R["main"]; SW, RG, MI = "Sliding window", "Plain RAG (FIFO)", "Memory-inspired"
NAVY, GREEN, GREY = RGBColor(0x1F, 0x2D, 0x4D), RGBColor(0x00, 0x7A, 0x5A), RGBColor(0x55, 0x55, 0x55)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def tb(sl, x, y, w, h, text, size=18, bold=False, color=None, align=None):
    box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.font.size = Pt(size); p.font.bold = bold; p.space_after = Pt(6)
        if color: p.font.color.rgb = color
    return box


def slide(title, bullets=None, fig=None, figw=7.2, note=None, size=23, bw=None):
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.0)); bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb(s, 0.5, 0.17, 12.3, 0.7, title, 30, True, RGBColor(255, 255, 255))
    if bullets:
        w = bw or (5.2 if fig else 12.3)
        tb(s, 0.5, 1.3, w, 5.5, ["• " + b for b in bullets], size)
    if fig:
        s.shapes.add_picture(str(FIG / fig), Inches((13.333 - figw) / 2 if figw > 9 else 13.0 - figw), Inches(1.3), width=Inches(figw))
    if note:
        tb(s, 0.5, 6.85, 12.3, 0.5, note, 12, False, GREY)
    return s


s = prs.slides.add_slide(blank)
tb(s, 0.8, 2.0, 11.5, 1.2, "A Memory-Inspired Chatbot", 48, True, NAVY)
tb(s, 0.8, 3.3, 11.5, 1.0, "Working memory, long-term memory and forgetting for dialogue", 26, False, GREEN)
tb(s, 0.8, 4.8, 11.5, 1.6, ["IT 207 Human Intelligence project", "[Team members]  |  [Instructor]  |  October 2026"], 18, False, GREY)

slide("Problem: chatbots forget, or remember everything",
      ["Stateless between sessions; finite context window",
       "Sliding window: old facts vanish abruptly",
       "Store-everything RAG: unbounded growth, noisy retrieval, stale and updated facts look equally valid",
       "Human memory is selective, bounded and adaptive",
       "**Question:** does a memory organised like human memory recall more under the same memory budget, and which parts matter?".replace("**", "")],
      note="Computational proxy, not a measurement of anyone's intelligence; inspired by, not a model of, the brain.")
slide("Theory behind the design",
      ["Atkinson & Shiffrin 1968: sensory, short-term, long-term stores; attention and rehearsal control transfer",
       "Baddeley & Hitch 1974; Miller 1956; Cowan 2001: limited-capacity working memory (~7, or ~4)",
       "Tulving 1972: episodic vs. semantic memory; McClelland et al. 1995: consolidation",
       "Ebbinghaus; Anderson & Schooler 1991; ACT-R: B = ln Σ (t − t_j)^−d, decay + rehearsal"])
slide("Architecture", fig="fig1_architecture.png", figw=10.0,
      note="Attention gate → working memory (7) → episodic + semantic long-term memory → activation-based forgetting under a budget.")
slide("Cognitive process → AI component",
      ["Selective attention → salience score + attention gate",
       "Working memory → 7-item buffer with displacement and rehearsal counting",
       "Episodic memory → raw utterances + BM25 retrieval",
       "Semantic memory → slot = value facts; newest value supersedes the old",
       "Forgetting / rehearsal → ACT-R activation; lowest-activation trace evicted over budget",
       "User control → /memory (inspect), /forget (delete); declines to answer about third parties"])
slide("Experimental design",
      [f"Synthetic multi-session personas: 5 sessions × 24 turns; facts, updates, third-party distractors, one-off episodes, chit-chat noise",
       f"Dev templates for tuning (30 personas); **held-out** test templates (100 personas, {R['n_test_queries']} questions)".replace("**", ""),
       "Question types: fact, update, unanswerable, episodic detail",
       "Baselines: stateless, sliding window, plain RAG (FIFO, k=3 and 10) at the same item budget; unbounded references",
       "Metrics: accuracy + 95% bootstrap CI, P/R/F1, hallucination, latency, footprint; LoCoMo evidence recall@k"])
slide(f"Result: {M[MI]['acc']:.2f} vs {M[SW]['acc']:.2f} vs {M[RG]['acc']:.2f} at budget 40", fig="fig4_by_type.png", figw=7.0,
      bullets=[f"Memory-inspired {M[MI]['acc']:.3f} (95% CI {M[MI]['ci'][0]:.3f}-{M[MI]['ci'][1]:.3f})",
               f"Window {M[SW]['acc']:.3f}; plain RAG {M[RG]['acc']:.3f}",
               f"Paired gain over window +{R['diff_vs_window'][0]:.2f} (CI excludes 0)",
               f"Full-context reference {M['Full context (ref., unbounded)']['acc']:.3f}: keeping everything is better when it fits"], bw=5.4, size=20)
slide("Budget sweep: help when memory is scarce", fig="fig3_budget_sweep.png", figw=6.6,
      bullets=["Advantage of +0.19 to +0.24 at budgets 12-40",
               "Flat after 40: it has already kept what it will keep",
               "At budget 120 the window wins (0.754 vs 0.676)",
               "Claim is narrow: selective memory under scarcity"], bw=5.8, size=20)
slide("Ablations: what actually matters", fig="fig5_ablation.png", figw=7.6,
      bullets=["Semantic memory: biggest effect (−0.06 / −0.10)",
               "Attention gate: −0.10 at B=40",
               "Decay: −0.05 at B=20, none at B=40",
               "Supersession: small",
               "Working memory: no benefit, slightly hurts at B=20 (negative result)"], bw=4.7, size=20)
slide("Perception vs. memory", fig="fig6_oracle.png", figw=6.5,
      bullets=[f"Extractor on held-out phrasing: P {R['extractor_test']['p']:.2f}, R {R['extractor_test']['r']:.2f}",
               f"Question parser covers {R['query_cov_test']*100:.0f}% of held-out questions",
               f"Oracle perception: {R['oracle']['40'][MI]['acc']:.2f} vs {R['oracle']['40'][SW]['acc']:.2f} (window) at B=40",
               "Bottleneck = rule-based perception → use an LLM extractor"], bw=5.9, size=20)
L = R["locomo"]
slide("Real data: LoCoMo evidence retrieval", fig="fig9_locomo.png", figw=6.2,
      bullets=[f"{R['locomo_stats']['qa']} questions, {R['locomo_stats']['convs']} human-reviewed conversations",
               f"Fixed budgets: higher recall@5 than BM25+FIFO ({L['100']['Memory-inspired']['recall']['5']:.3f} vs {L['100']['BM25 (FIFO)']['recall']['5']:.3f} at 100 turns)",
               f"Unbounded: BM25 is better ({L['None']['BM25 (FIFO)']['recall']['5']:.3f} vs {L['None']['Memory-inspired']['recall']['5']:.3f})",
               "Retrieval only: no reader evaluated"], bw=6.0, size=20)
demo = (ROOT / "results" / "demo_transcript.txt").read_text().split("\n")
pick = [l for l in demo if l.startswith(("you> What is my name", "you> Where do I live", "you> Where does my sister", "you> What is my favorite color", "you> /forget", "bot> Your name", "bot> You live", "bot> I only", "bot> I don't know", "bot> Forgot"))][:10]
s = slide("Live demo: persistence, update, abstention, control", ["python chat.py   (/memory, /forget, /skip)"], size=18)
tb(s, 0.5, 2.1, 12.3, 4.5, pick, 20)
slide("Ethics and responsible use",
      ["Privacy: bounded retention; /memory transparency; /forget deletion; no answers about third parties",
       "Gaps: third-party text still stored until evicted; no encryption or user separation",
       "Bias: English, capitalised Latin-script names; heuristic salience decides what is 'important'",
       "No over-claiming: not a measure of intelligence; ACT-R values borrowed, not fitted"])
slide("Limitations and future work",
      ["Synthetic templated benchmark; rule-based perception; templated replies; no human evaluation",
       "Working memory buffer did not help here; budget counted in items, not tokens",
       "Next: LLM extractor/reader; embeddings; learned salience; LongMemEval and MSC; consolidation by summarising; privacy engineering"])
slide("Conclusions",
      [f"Under scarce memory, a cognitively organised memory answered {M[MI]['acc']:.2f} vs {M[SW]['acc']:.2f}/{M[RG]['acc']:.2f}, and {R['oracle']['40'][MI]['acc']:.2f} with perfect perception",
       "Gain comes from compact semantic memory and selective encoding; decay matters when memory is tight",
       "Not better than keeping everything when everything fits",
       "Theory-inspired design is useful for deciding what to forget, not automatically smarter",
       "Code, data generator and all numbers are reproducible (see README)"])
prs.save(ROOT / "report" / "Memory_Inspired_Chatbot_Slides.pptx")
print("slides:", len(prs.slides))
